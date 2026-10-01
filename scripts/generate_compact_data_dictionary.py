import docx
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_ALIGN_VERTICAL
from docx.oxml import parse_xml, OxmlElement
from docx.oxml.ns import nsdecls, qn

def set_cell_background(cell, fill_hex):
    tcPr = cell._tc.get_or_add_tcPr()
    shd = parse_xml(f'<w:shd {nsdecls("w")} w:fill="{fill_hex}"/>')
    tcPr.append(shd)

def set_cell_margins(cell, top=80, bottom=80, left=100, right=100):
    tcPr = cell._tc.get_or_add_tcPr()
    tcMar = parse_xml(f'<w:tcMar {nsdecls("w")}><w:top w:w="{top}" w:type="dxa"/><w:bottom w:w="{bottom}" w:type="dxa"/><w:left w:w="{left}" w:type="dxa"/><w:right w:w="{right}" w:type="dxa"/></w:tcMar>')
    tcPr.append(tcMar)

def set_table_borders(table, color="CCCCCC", sz="4", val="single"):
    tblPr = table._tbl.tblPr
    borders = parse_xml(
        f'<w:tblBorders {nsdecls("w")}>'
        f'<w:top w:val="{val}" w:sz="{sz}" w:space="0" w:color="{color}"/>'
        f'<w:bottom w:val="{val}" w:sz="{sz}" w:space="0" w:color="{color}"/>'
        f'<w:left w:val="{val}" w:sz="{sz}" w:space="0" w:color="{color}"/>'
        f'<w:right w:val="{val}" w:sz="{sz}" w:space="0" w:color="{color}"/>'
        f'<w:insideH w:val="{val}" w:sz="{sz}" w:space="0" w:color="{color}"/>'
        f'<w:insideV w:val="{val}" w:sz="{sz}" w:space="0" w:color="{color}"/>'
        f'</w:tblBorders>'
    )
    tblPr.append(borders)

def format_cell(cell, width, bg_color=None, font_name="Courier New", font_size=8.5, bold=False, text_color=RGBColor(30, 41, 59), align=WD_ALIGN_PARAGRAPH.LEFT):
    cell.width = width
    if bg_color:
        set_cell_background(cell, bg_color)
    set_cell_margins(cell, top=70, bottom=70, left=90, right=90)
    for p in cell.paragraphs:
        p.alignment = align
        p.paragraph_format.space_before = Pt(0)
        p.paragraph_format.space_after = Pt(0)
        p.paragraph_format.line_spacing = 1.05
        for r in p.runs:
            r.font.name = font_name
            r.font.size = Pt(font_size)
            r.font.bold = bold
            r.font.color.rgb = text_color

def build_compact_data_dictionary():
    doc = docx.Document()

    # Narrow margins (0.5 to 0.6 inch) for compact, clean page budget
    for section in doc.sections:
        section.top_margin = Inches(0.6)
        section.bottom_margin = Inches(0.6)
        section.left_margin = Inches(0.6)
        section.right_margin = Inches(0.6)

    # Document Title
    p_title = doc.add_paragraph()
    p_title.paragraph_format.space_before = Pt(0)
    p_title.paragraph_format.space_after = Pt(2)
    r_title = p_title.add_run("Data Dictionary")
    r_title.font.name = "Courier New"
    r_title.font.size = Pt(16)
    r_title.font.bold = True
    r_title.font.color.rgb = RGBColor(20, 50, 30)

    # Subtitle / Intro paragraph
    p_intro = doc.add_paragraph()
    p_intro.paragraph_format.space_before = Pt(0)
    p_intro.paragraph_format.space_after = Pt(8)
    p_intro.paragraph_format.line_spacing = 1.15
    r_intro = p_intro.add_run(
        "Updated based on the revised 12-phase system process flow. The data entities below support account approval, "
        "farm setup, cropping, field monitoring, AI processing, expert validation, GIS monitoring, management actions, "
        "report resolution, LGU-MAO monitoring, harvest recording, and cycle completion across the VISAIA mobile app and web dashboard."
    )
    r_intro.font.name = "Courier New"
    r_intro.font.size = Pt(8.5)
    r_intro.font.color.rgb = RGBColor(50, 60, 50)

    # Data Rows
    table_data = [
        (
            "USERS",
            "Stores all system user accounts across mobile app and web dashboards, including farmers, MAO district administrators, RCPC regional specialists, and system admins.",
            "uid (PK), email, name/fullName, role (farmer/mao_admin/rcpc_admin), userType (app/dashboard), isVerified, status (pending/approved/rejected), createdAt",
            "FARMERS, VALIDATION_REQUESTS, EXPERT_VALIDATION, ALERTS, ALERT_RECIPIENTS"
        ),
        (
            "FARMERS",
            "Stores farmer agricultural profiles, demographic attributes, RSBSA registration, declared farm size, and verification documents.",
            "id/uid (PK/FK), fullName, sex, birthdate, farmSize, hasRsbsaId, rsbsaId, skipFarmerId, farmerIdImage, status (pending/verified/rejected), verifiedBy (FK), verifiedAt, rejectionReason, createdAt",
            "USERS, VALIDATION_REQUESTS, FARM_LOCATIONS, CROPPING_CYCLES, PEST_REPORTS"
        ),
        (
            "VALIDATION_REQUESTS",
            "Manages farmer account approval and registration verification requests submitted for LGU-MAO administrative review and audit trail.",
            "id (PK), farmerId (FK), status (pending/approved/rejected), reviewedBy (FK), reviewedAt, reviewNotes/rejectionReason, createdAt",
            "FARMERS, USERS (MAO Admin)"
        ),
        (
            "FARM_LOCATIONS",
            "Stores the geospatial boundary polygons, total acreage, centroid GPS coordinates, and administrative district for each farm parcel.",
            "id (PK), farmerId (FK), name, acres, hectares, boundaries (Array<LatLng>), location (GeoPoint), district, createdAt",
            "FARMERS, FARM_FIELDS, RISK_ZONES, HISTORICAL_INVESTIGATIONS"
        ),
        (
            "FARM_FIELDS",
            "Stores individual field plot subdivisions within a farm parcel, specific boundary polygons, acreage, and assigned crops.",
            "id (PK), farmId (FK), farmerId (FK), name, acres, crop, boundaries (Array<LatLng>), createdAt",
            "FARM_LOCATIONS, CROPPING_CYCLES, PEST_REPORTS, TRAP_MONITORING"
        ),
        (
            "CROPPING_CYCLES",
            "Tracks seasonal crop growing cycles, botanical varieties, phenological growth stages, harvest schedules, and completion states.",
            "id (PK), farmId (FK), fieldId (FK), cycleName, cropVariety, plantingDate, harvestDate, actualHarvestDate, seedDensity, growthStage (SEEDLING..MATURITY), isCompleted, status (growing/harvested/completed), netIncome, grossIncome, createdAt",
            "FARM_LOCATIONS, FARM_FIELDS, PLANTING_RECORDS, FIELD_MONITORING_TASKS, PEST_REPORTS, HARVEST_RECORDS, FINANCIAL_OUTCOMES"
        ),
        (
            "PLANTING_RECORDS",
            "Stores botanical details and planting specifications recorded when initializing a new cycle (seed variety, density, sowing date, spacing).",
            "id (PK), cycleId (FK), plantingDate, cropVariety/seedType, seedDensity/quantity, sowingMethod, spacingRowCm, spacingHillCm",
            "CROPPING_CYCLES"
        ),
        (
            "FIELD_MONITORING_TASKS",
            "Tracks weekly phenological milestone timelines, Days After Planting (DAP), scouting completion, and scheduled field tasks.",
            "id/weekId (PK), cycleId (FK), weekNumber, title, growthStage, startDate, endDate, scoutingStatus (pending/completed), fawStatus, tasks",
            "CROPPING_CYCLES, DAILY_ACTIVITY_LOGS, PEST_REPORTS"
        ),
        (
            "DAILY_ACTIVITY_LOGS",
            "Stores granular day-to-day farm operation logs (irrigation, fertilization, pesticide/bio-spray application, weeding, and scouting).",
            "id (PK), cycleId (FK), fieldId (FK), type (irrigation/fertilization/pesticide/weeding/scouting), title, description/notes, date, dayNumber/DAP, status, images, timestamp",
            "CROPPING_CYCLES, FARM_FIELDS, FIELD_MONITORING_TASKS"
        ),
        (
            "TRAP_MONITORING",
            "Stores pheromone and light trap inspection logs, adult FAW moth counts, lure condition, inspection photos, and date.",
            "id (PK), cycleId (FK), fieldId (FK), farmerId (FK), trapName, trapType (Pheromone/Light/Sticky), mothCount, condition, monitoringDate, notes, images, timestamp",
            "FARM_FIELDS, CROPPING_CYCLES, PEST_REPORTS, SPATIAL_CLUSTERS, MANAGEMENT_ACTIONS"
        ),
        (
            "PEST_REPORTS",
            "Stores field scouting and AI camera detection submissions, photos, coordinates, pest counts, severity, risk level, and resolution status.",
            "id (PK), farmerId (FK), farmId (FK), fieldId (FK), cycleId (FK), reportType (regular/clustered), detection, scientificName, lifeStage, cropAffected, severity (low/moderate/high/critical), riskLevel, confidence, boxes, analysis, treatment, imageUrl, annotatedImageUrl, location (GeoPoint), status (pending/validated/rejected/resolved), resolutionExplanation, timestamp",
            "FARMERS, CROPPING_CYCLES, FARM_FIELDS, AI_PROCESSING, AI_DETECTIONS, VALIDATION_QUEUE, RISK_ZONE_SOURCE, SPATIAL_CLUSTERS, MANAGEMENT_ACTIONS, REPORT_RESOLUTION, ALERTS"
        ),
        (
            "AI_PROCESSING",
            "Stores YOLO AI inference job execution metadata, model version, endpoint latency, and execution status.",
            "id (PK), reportId (FK), modelVersion (YOLOv8-FAW), processingTimeMs, status (success/failed), endpoint, processedAt",
            "PEST_REPORTS, AI_DETECTIONS"
        ),
        (
            "AI_DETECTIONS",
            "Stores individual pest object detections from YOLO model inference with bounding boxes, confidence scores, and instar classifications.",
            "id (PK), reportId (FK), pestType/className, confidence, x, y, width, height, lifeStage",
            "PEST_REPORTS, AI_PROCESSING, EXPERT_VALIDATION, MITIGATION_ADVISORIES"
        ),
        (
            "VALIDATION_QUEUE",
            "Manages pest reports awaiting examination, triage, and diagnosis confirmation by RCPC crop protection specialists.",
            "id (PK), reportId (FK), submittedBy (FK), submittedDate, assignedExpert (FK), priority (urgent/high/normal), status (pending/under_review/validated/rejected)",
            "PEST_REPORTS, USERS (RCPC Specialist), EXPERT_VALIDATION"
        ),
        (
            "EXPERT_VALIDATION",
            "Stores official RCPC expert review verdicts, taxonomic confirmations, diagnostic corrections, risk assessments, and justifications.",
            "id (PK), reportId (FK), reviewerId (FK), reviewerName, validationStatus (confirmed/corrected/rejected/false_positive), validatedPest, validatedLifeStage, validatedSeverity, expertNotes/comments, rejectionReason, validatedAt",
            "PEST_REPORTS, VALIDATION_QUEUE, USERS (RCPC Reviewer), MITIGATION_ADVISORIES"
        ),
        (
            "MITIGATION_ADVISORIES",
            "Stores tailored pest control advisories, chemical/bio safety protocols, and IPM management directives sent to farmers after validation.",
            "id (PK), reportId (FK), farmerId (FK), advisoryText, controlMethods, urgencyLevel (immediate/soon/preventive), sentAt",
            "PEST_REPORTS, EXPERT_VALIDATION, USERS (Farmers), MANAGEMENT_ACTIONS"
        ),
        (
            "MANAGEMENT_ACTIONS",
            "Stores farmer-executed pest mitigation interventions (biological, cultural, or chemical spray), product used, dosage, and date.",
            "id (PK), farmerId (FK), reportId (FK), controlMethod (Biological/Cultural/Chemical), chemicalName, dosageRate, actionDate, status (completed/in_progress), notes",
            "FARMERS, PEST_REPORTS, DAILY_ACTIVITY_LOGS, REPORT_RESOLUTION"
        ),
        (
            "REPORT_RESOLUTION",
            "Captures formal pest threat resolution documentation, post-mitigation field observations, clearance explanation, and status update.",
            "id/reportId (PK), reportId (FK), resolvedBy (FK), resolvedByUserName, resolutionExplanation, resolvedAt, status (resolved)",
            "PEST_REPORTS, FARMERS, USERS"
        ),
        (
            "RISK_ZONES",
            "Geospatial outbreak containment zones and contagion buffer polygons calculated from pest concentrations and wind trajectories.",
            "id (PK), name, riskLevel (Low/Medium/High/Critical), zoneGeometry (GeoJSON Polygon / Center + Radius), affectedFarmCount, activeReportCount, createdAt",
            "RISK_ZONE_SOURCE, SPATIAL_CLUSTERS, ALERTS, FARM_LOCATIONS, REGIONAL_ADVISORY_ZONES"
        ),
        (
            "RISK_ZONE_SOURCE",
            "Junction tracking individual scout reports and trap catches that contribute statistical weight to each geospatial risk zone.",
            "id (PK), riskZoneId (FK), reportId (FK), contributionWeight, distanceFromCenterKm",
            "RISK_ZONES, PEST_REPORTS, SPATIAL_CLUSTERS"
        ),
        (
            "SPATIAL_CLUSTERS",
            "Aggregated spatial cluster nodes generated from synchronized field scouting batches, moth trap counts, and plant damage rates.",
            "id (PK), cycleId (FK), reportType (clustered), dap, growthStage, totalDamaged, totalInspected, damagePercentage, totals (eggs/larvae/pupae/moths), timestamp",
            "CROPPING_CYCLES, PEST_REPORTS, RISK_ZONES, RISK_ZONE_SOURCE"
        ),
        (
            "ALERTS",
            "Broadcast alerts and outbreak warnings published by RCPC/MAO for targeted districts, specifying severity, affected crops, and IPM measures.",
            "id (PK), createdBy (FK), senderName, senderRole (rcpc_admin/mao_admin), title, description/message, pestSpecies, affectedCrops, severity (info/warning/critical), targetAudience, affectedRegions, affectedFarms, recommendations, status, publishedAt, validUntil, viewCount",
            "RISK_ZONES, ALERT_RECIPIENTS, PEST_REPORTS, USERS"
        ),
        (
            "ALERT_RECIPIENTS",
            "Tracks in-app notifications and delivery receipts dispatched to targeted farmers (outbreak alerts, proximity warnings, verification updates).",
            "id (PK), alertId (FK), userId/farmerId (FK), type (report_pending/report_proximity/farmer_registered/advisory), title, message, read (Boolean), data, createdAt",
            "ALERTS, USERS (Farmers)"
        ),
        (
            "REGIONAL_ADVISORIES",
            "Broad agro-ecological bulletins and seasonal pest advisories published by RCPC for provinces and regional agricultural districts.",
            "id (PK), issuedBy/createdBy (FK), title, content/description, severityLevel (info/warning/critical), validFrom, validTo, createdAt",
            "USERS (RCPC Admin), REGIONAL_ADVISORY_ZONES, RISK_ZONES"
        ),
        (
            "REGIONAL_ADVISORY_ZONES",
            "Junction linking regional advisory bulletins to specific administrative districts, municipalities, and active risk zones.",
            "id (PK), advisoryId (FK), riskZoneId/regionName (FK)",
            "REGIONAL_ADVISORIES, RISK_ZONES"
        ),
        (
            "PEST_REFERENCES",
            "Master taxonomic reference catalog for monitored agricultural pests (common name, scientific name, biology, host crops, economic threshold).",
            "id/slug (PK), commonName, scientificName, description, hostCrops, damageType, economicThreshold, integratedManagement, images, updatedAt",
            "PEST_LIFESTAGE_INFO, IPM_RECOMMENDATIONS, PEST_REPORTS"
        ),
        (
            "PEST_LIFESTAGE_INFO",
            "Detailed reference data on biological life cycle stages (instar durations, morphological identification features, and crop damage potential).",
            "id (PK), pestId/slug (FK), lifeStage (Egg/Early Larva/Late Larva/Pupa/Adult), durationDays, damagePotential (Critical/High/Moderate/None), distinguishingFeatures",
            "PEST_REFERENCES"
        ),
        (
            "IPM_RECOMMENDATIONS",
            "Integrated Pest Management protocols detailing cultural, biological, and chemical control methods, pesticide active ingredients, and pre-harvest intervals.",
            "id (PK), pestId/slug (FK), controlMethod (Cultural/Biological/Chemical), recommendationText, activeChemical, rateApplication, preharvestIntervalDays",
            "PEST_REFERENCES, MITIGATION_ADVISORIES"
        ),
        (
            "HISTORICAL_INVESTIGATIONS",
            "Archive of historical pest outbreak cases, peak damage percentages, weather correlations, and evaluated mitigation efficacy.",
            "id (PK), locationId/farmId (FK), pestType, eventDate, peakInfestationRate, successfulMitigation, recordedBy (FK)",
            "FARM_LOCATIONS, USERS, CROPPING_CYCLES"
        ),
        (
            "HARVEST_RECORDS",
            "Captures end-of-cycle crop harvest data: actual vs expected date, total yield, good vs damaged yield, loss rate, unit market price, and early harvest reason.",
            "id/cycleId (PK/FK), actualHarvestDate, expectedHarvestDate, actualYield, goodYield, damagedYield, lossRate (%), marketPrice (PHP), totalValue/grossIncome (PHP), isEarlyHarvest, earlyHarvestReason, harvestNotes, status (harvested)",
            "CROPPING_CYCLES, FINANCIAL_OUTCOMES"
        ),
        (
            "FINANCIAL_OUTCOMES",
            "Calculates comprehensive financial outcome of completed cycle: gross revenue, breakdown of input expenses (seeds, fertilizer, bio-spray/chemicals, labor), pest losses, and net profit.",
            "id/cycleId (PK/FK), grossIncome (PHP), seedsCost (PHP), fertilizerCost (PHP), controlMethodCost (PHP), otherExpenses (PHP), totalCycleCost (PHP), damageCost/pestLoss (PHP), netIncome/income (PHP), isAuditCompleted, auditCompletedAt",
            "CROPPING_CYCLES, HARVEST_RECORDS"
        ),
        (
            "SYSTEM_CONFIG",
            "Remote configuration controlling mobile app version gating, minimum build requirements, download URLs, and system maintenance mode.",
            "id (PK: 'config'), latestVersion, latestBuild, minVersion, minBuild, downloadUrl, updateMessage, maintenanceMode (Boolean)",
            "USERS (Mobile App Clients)"
        )
    ]

    # Create 4-column table
    col_widths = [Inches(1.5), Inches(2.3), Inches(2.2), Inches(1.3)]
    table = doc.add_table(rows=len(table_data) + 1, cols=4)
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.autofit = False
    set_table_borders(table, color="B0BEC5", sz="4", val="single")

    # Header Row
    headers = ["Table Name", "Description", "Key Fields", "Connect To"]
    header_row = table.rows[0]
    for idx, h_text in enumerate(headers):
        cell = header_row.cells[idx]
        cell.text = h_text
        format_cell(cell, col_widths[idx], bg_color="1E3A1E", font_name="Courier New", font_size=9, bold=True, text_color=RGBColor(255, 255, 255))

    # Data Rows
    for r_idx, (t_name, desc, key_fields, connect_to) in enumerate(table_data, start=1):
        row = table.rows[r_idx]
        is_even = (r_idx % 2 == 0)
        bg = "F4F7F4" if is_even else "FFFFFF"

        # Table Name Cell
        row.cells[0].text = t_name
        format_cell(row.cells[0], col_widths[0], bg_color=bg, font_name="Courier New", font_size=8.5, bold=True, text_color=RGBColor(20, 60, 30))

        # Description Cell
        row.cells[1].text = desc
        format_cell(row.cells[1], col_widths[1], bg_color=bg, font_name="Courier New", font_size=8, bold=False, text_color=RGBColor(30, 41, 59))

        # Key Fields Cell
        row.cells[2].text = key_fields
        format_cell(row.cells[2], col_widths[2], bg_color=bg, font_name="Courier New", font_size=7.5, bold=False, text_color=RGBColor(30, 41, 59))

        # Connect To Cell
        row.cells[3].text = connect_to
        format_cell(row.cells[3], col_widths[3], bg_color=bg, font_name="Courier New", font_size=7.5, bold=False, text_color=RGBColor(40, 70, 50))

    output_path = r'c:\PROJECTS\visaia\Updated_Data_Dictionary.docx'
    doc.save(output_path)
    print(f"Successfully generated compact Courier New Data Dictionary at: {output_path}")

if __name__ == '__main__':
    build_compact_data_dictionary()
